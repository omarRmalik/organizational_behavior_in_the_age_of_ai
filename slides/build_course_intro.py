"""Build slides/course_intro.pptx: an introduction to the course covering Parts 0-IV."""
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = sys.argv[1] if len(sys.argv) > 1 else "course_intro.pptx"

# Palette: teal = knowing, coral = deciding, gold = consequences; deep slate dominates.
DARK = RGBColor(0x1F, 0x2A, 0x2E)
TEAL = RGBColor(0x2E, 0x8B, 0x7A)
CORAL = RGBColor(0xD9, 0x69, 0x4A)
GOLD = RGBColor(0xC9, 0x8A, 0x1E)
TINT = RGBColor(0xEE, 0xF3, 0xF1)
MUTED = RGBColor(0x5B, 0x6B, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_TEAL = RGBColor(0xA8, 0xD5, 0xCB)
HEAD = "Cambria"
BODY = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------
def text(slide, x, y, w, h, paras, size=16, color=DARK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, space_after=0, margin=0):
    """paras: str, or list of str / list of (text, {overrides}) runs per paragraph."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for run_text, o in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def shape(slide, kind, x, y, w, h, fill, line=None, radius=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if radius is not None:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, fill=TINT):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.06)


def badge(slide, x, y, d, fill, label, size=16, color=WHITE):
    c = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = BODY, Pt(size), True, color
    return c


def motif(slide, x, y, d, on_dark=False):
    """Two overlapping circles: the human and the AI, knowing and deciding together."""
    a = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, CORAL)
    b = shape(slide, MSO_SHAPE.OVAL, x + d * 0.62, y, d, d, TEAL)
    # python-pptx has no transparency setting; add an alpha element so the overlap shows
    for s in (a, b):
        sf = s.fill._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        alpha = sf.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", {"val": "78000"})
        sf.append(alpha)


def new_slide(part=None, title=None, dark=False, notes=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = DARK if dark else WHITE
    if part:
        text(s, 0.6, 0.42, 9, 0.3, part.upper(), size=11, bold=True, color=LIGHT_TEAL if dark else TEAL)
    if title:
        text(s, 0.6, 0.72, 11.2, 0.9, title, size=32, font=HEAD, bold=True, color=WHITE if dark else DARK)
        motif(s, 12.05, 0.55, 0.42)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def table(slide, x, y, w, col_w, rows, header_fill=DARK, size=12, row_h=0.4, first_col_bold=False):
    t = slide.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (TINT if r % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = val
            run.font.name = BODY
            run.font.size = Pt(size)
            run.font.bold = r == 0 or (first_col_bold and c == 0)
            run.font.color.rgb = WHITE if r == 0 else DARK
    return t


def arrow(slide, x, y, w=0.45, h=0.4, fill=MUTED):
    return shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, fill)


# ---------- 1. Title ----------
s = new_slide(dark=True, notes="Welcome to MGMT-205. This term we study organizational behavior in a setting where "
              "people increasingly work alongside AI. The two overlapping circles that appear throughout the deck "
              "represent the course's core idea: humans and AI knowing and deciding together.")
text(s, 0.8, 1.2, 7, 0.4, "MGMT-205  ·  FALL 2026", size=14, bold=True, color=LIGHT_TEAL)
text(s, 0.8, 1.75, 7.4, 2.6, "Organizational Behavior in the Age of Artificial Intelligence", size=44, font=HEAD,
     bold=True, color=WHITE)
text(s, 0.8, 4.45, 7.2, 1.0, "How organizational behavior changes when humans and AI jointly develop knowledge "
     "and make decisions", size=20, color=LIGHT_TEAL, italic=True)
text(s, 0.8, 6.3, 8, 0.4, "Omar R. Malik, Ph.D.  ·  School of Management  ·  Kettering University", size=14,
     color=WHITE)
motif(s, 8.9, 2.1, 2.7)
text(s, 8.9, 4.95, 2.7, 0.4, "HUMAN", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 10.57, 4.95, 2.7, 0.4, "AI", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ---------- 2. Why this course exists ----------
s = new_slide("Part 0 · Course Philosophy", "Why this course exists",
              notes="Classic OB theories were built for workplaces where only people held knowledge and only people "
                    "made decisions. Walk through the four vignettes: in each, knowledge or a decision is now shared "
                    "between a person and an AI system.")
text(s, 0.6, 1.95, 4.6, 4.5, [
    [("Classic OB was built for workplaces where ", {}), ("only people held knowledge", {"bold": True}),
     (" and ", {}), ("only people made decisions.", {"bold": True})],
    [("That assumption no longer holds.", {"color": CORAL, "bold": True})],
    "In each of these moments, knowledge is built jointly and decisions are made jointly. OB must account for what "
    "happens in those moments.",
], size=18, space_after=14)
vignettes = [
    ("An employee drafts a report with a generative AI assistant", TEAL),
    ("A recruiter screens applicants with a predictive algorithm", CORAL),
    ("A team debates a recommendation no one can fully explain", TEAL),
    ("A professional explains to a client a decision software made", CORAL),
]
for i, (v, col) in enumerate(vignettes):
    cx, cy = 5.7 + (i % 2) * 3.65, 1.95 + (i // 2) * 2.35
    card(s, cx, cy, 3.4, 2.05)
    badge(s, cx + 0.3, cy + 0.3, 0.55, col, str(i + 1))
    text(s, cx + 0.3, cy + 1.0, 2.85, 0.95, v, size=15, color=DARK)

# ---------- 3. Central question ----------
s = new_slide("Part 0 · The Central Question", dark=True,
              notes="Everything in the course returns to this question. Traditional OB asks what motivates people, "
                    "how groups decide, what makes leaders effective. We keep those questions and add a second one "
                    "to every topic: what does this theory assume about who knows and who decides?")
motif(s, 5.85, 1.0, 0.95)
text(s, 1.2, 2.3, 10.9, 2.2, "How does organizational behavior change when humans and AI jointly develop knowledge "
     "and make decisions?", size=36, font=HEAD, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 1.8, 4.9, 9.7, 1.2, "And for every topic: what does this theory assume about who knows and who decides — "
     "and what happens to it when AI becomes part of knowing and deciding?", size=18, italic=True, color=LIGHT_TEAL,
     align=PP_ALIGN.CENTER)

# ---------- 4. Two processes ----------
s = new_slide("Part 0 · Course Philosophy", "Two processes at the heart of organizational life",
              notes="Almost everything OB studies depends on knowing and deciding. AI now participates in both. "
                    "The JKD Framework in Part IV is built on these two processes plus their consequences.")
for i, (letter, col, head, body, classic) in enumerate([
    ("K", TEAL, "Knowing", "What employees perceive, learn, remember, and share.",
     "Classic view: knowledge is created by people through experience and conversation (Nonaka, 1994), within "
     "the limits of bounded rationality (Simon, 1991)."),
    ("D", CORAL, "Deciding", "Who is hired, how work is designed, which option a team chooses, how performance "
     "is judged.", "Classic view: decisions are human judgments shaped by personality, perception, group "
     "dynamics, and leadership."),
]):
    cx = 0.6 + i * 6.2
    card(s, cx, 1.95, 5.9, 3.55)
    badge(s, cx + 0.35, 2.25, 0.8, col, letter, size=24)
    text(s, cx + 1.35, 2.35, 4.2, 0.6, head, size=26, font=HEAD, bold=True, color=col)
    text(s, cx + 0.35, 3.25, 5.2, 2.1, [body, [(classic, {"color": MUTED, "size": 14})]], size=17, space_after=10)
text(s, 0.6, 5.85, 12.1, 0.9, [[("AI now participates in both, ", {"bold": True}),
     ("changing trust, learning, job satisfaction, leadership, team decisions, fairness, and professional "
      "expertise.", {})]], size=18)

# ---------- 5. Context vs content; three sources ----------
s = new_slide("Part 0 · Course Philosophy", "AI is the context. OB is the content.",
              notes="This is not a course about how AI works and requires no programming. Each week from Week 2 "
                    "places three sources side by side. The goal is not to prove classic theory wrong; often it "
                    "holds up well.")
steps = [("1", "Textbook chapter", "Core OB concepts", DARK),
         ("2", "Classic article", "The foundational theory", TEAL),
         ("3", "Contemporary AI article", "JOB, AMJ, ASQ, JMS, Personnel Psychology", CORAL)]
for i, (n, head, sub, col) in enumerate(steps):
    cx = 0.6 + i * 4.2
    card(s, cx, 2.0, 3.6, 2.2)
    badge(s, cx + 0.3, 2.3, 0.6, col, n)
    text(s, cx + 0.3, 3.05, 3.1, 0.5, head, size=19, bold=True)
    text(s, cx + 0.3, 3.5, 3.1, 0.6, sub, size=14, color=MUTED)
    if i < 2:
        arrow(s, cx + 3.68, 2.9)
card(s, 0.6, 4.6, 12.1, 2.1, fill=WHITE).line.color.rgb = LIGHT_TEAL
text(s, 0.95, 4.8, 11.4, 1.8, [
    [("The goal: ", {"bold": True, "color": TEAL}),
     ("see exactly where a classic theory holds, where it bends, and where it needs new ideas.", {})],
    [("Example: ", {"bold": True, "color": CORAL}),
     ("Schulz et al. (2025) use job characteristics theory — nearly fifty years old — to explain why AI adoption "
      "both raises and lowers job satisfaction.", {})],
], size=17, space_after=12)

# ---------- 6. Assumptions table ----------
s = new_slide("Part I · Course Foundations", "What AI changes in classic OB assumptions",
              notes="Each week of the course examines one row of this table in depth. Ask students which row "
                    "surprises them most.")
table(s, 0.6, 1.9, 12.1, [5.6, 6.5], [
    ["Classic assumption", "What AI changes"],
    ["Knowledge is created by people", "AI generates, organizes, and recombines knowledge"],
    ["Decisions are made by people", "AI recommends, ranks, and sometimes decides"],
    ["Trust is built between people", "Employees calibrate trust in systems they cannot inspect"],
    ["Managers design jobs that motivate", "Algorithms assign, monitor, and evaluate work"],
    ["Leaders influence through relationships", "Leaders collaborate and share influence with AI"],
    ["Groups fail when information isn't shared", "AI can pool information and create new asymmetries"],
    ["Fairness depends on managers' judgment", "Fairness depends on data, design, and how outputs are used"],
    ["Expertise gives professionals authority", "AI challenges who holds expertise and authority"],
], size=15, row_h=0.53)

# ---------- 7. Educational philosophy ----------
s = new_slide("Part I · Course Foundations", "Five commitments behind the course",
              notes="These five commitments shape how every week runs: we learn the foundation first, look for the "
                    "assumptions, rely on evidence, write to sharpen thinking, and keep people at the center.")
principles = [("The foundations still matter", "You can't evaluate how AI changes OB without first understanding OB."),
              ("Assumptions are where theories break", "Look for what a theory assumes about who knows and who decides."),
              ("Evidence beats opinion", "Rely on peer-reviewed research, not predictions about AI."),
              ("Writing deepens understanding", "Weekly reflections turn concepts into specific theory updates."),
              ("People come first", "Effective, fair, and humane work is the goal; AI is a means.")]
cols = [TEAL, CORAL, GOLD, TEAL, CORAL]
for i, ((head, body), col) in enumerate(zip(principles, cols)):
    cx = 0.6 + i * 2.46
    card(s, cx, 2.0, 2.26, 4.6)
    badge(s, cx + 0.25, 2.3, 0.6, col, str(i + 1))
    text(s, cx + 0.25, 3.1, 1.85, 1.2, head, size=17, bold=True)
    text(s, cx + 0.25, 4.4, 1.85, 2.0, body, size=14, color=MUTED)

# ---------- 8. Weekly cycle ----------
s = new_slide("Part I · Pedagogical Model", "The weekly cycle",
              notes="Week 1 introduces the course. From Week 2 on, every week follows this three-part cycle.")
cycle = [("Learn", "Read the textbook chapter and the classic article; review core concepts in class.", DARK),
         ("Question", "Read the AI article; identify which assumptions it supports, challenges, or changes.", TEAL),
         ("Apply", "Discuss an HBS case, write a reading reflection, and apply the ideas to the Work Study.", CORAL)]
for i, (head, body, col) in enumerate(cycle):
    cx = 0.6 + i * 4.2
    c = shape(s, MSO_SHAPE.OVAL, cx + 0.95, 1.95, 1.7, 1.7, col)
    tf = c.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = head
    tf.margin_left = tf.margin_right = 0
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = HEAD, Pt(20), True, WHITE
    text(s, cx, 3.95, 3.6, 1.5, body, size=17, align=PP_ALIGN.CENTER)
    if i < 2:
        arrow(s, cx + 3.65, 2.6)
text(s, 0.6, 6.0, 12.1, 0.6, "Week 1 introduces the course and the JKD Framework. Paired readings begin in Week 2.",
     size=15, italic=True, color=MUTED, align=PP_ALIGN.CENTER)

# ---------- 9. Signature learning experiences ----------
s = new_slide("Part I · Course Foundations", "Signature learning experiences",
              notes="Three experiences carry most of the learning: the weekly reading reflections with their OB "
                    "Theory Updates, the team Human–AI Work Study, and seven HBS case discussions.")
experiences = [("9", "Weekly reading reflections", "One each week from Week 2 to 10, ending with an OB Theory "
                "Update. Best 8 of 9 count.", TEAL),
               ("3", "Levels in the Human–AI Work Study", "Teams of three analyze a real organization's human–AI "
                "work at the individual, group, and organization levels.", CORAL),
               ("7", "Harvard Business School cases", "From Microsoft's Copilot deployment to JPMorganChase's "
                "leadership in the age of GenAI.", GOLD)]
for i, (num, head, body, col) in enumerate(experiences):
    cx = 0.6 + i * 4.1
    card(s, cx, 1.95, 3.8, 4.7)
    text(s, cx + 0.35, 2.15, 3.1, 1.3, num, size=72, font=HEAD, bold=True, color=col)
    text(s, cx + 0.35, 3.6, 3.1, 0.9, head, size=19, bold=True)
    text(s, cx + 0.35, 4.55, 3.1, 1.9, body, size=15, color=MUTED)

# ---------- 10. Expected mindset ----------
s = new_slide("Part I · Course Foundations", "The mindset that succeeds here",
              notes="Emphasize the second point: AI is a partner to question, not an authority to obey. Students' "
                    "own reactions to AI are themselves data about organizational behavior.")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.95, 5.2, 4.7, DARK, radius=0.06)
text(s, 1.0, 2.4, 4.4, 3.8, [[("“", {"size": 60, "color": CORAL, "font": HEAD})],
     "Treat AI as a partner to question, not an authority to obey."], size=28, font=HEAD, bold=True,
     color=WHITE)
mind = ["Come to class prepared to discuss the readings",
        "Notice your own reactions to AI — trust, doubt, relief, frustration — as data about OB",
        "Support every claim with evidence",
        "Treat classmates as collaborators in building knowledge"]
for i, m in enumerate(mind):
    cy = 2.0 + i * 1.18
    badge(s, 6.3, cy + 0.08, 0.5, [TEAL, CORAL, GOLD, TEAL][i], str(i + 1), size=14)
    text(s, 7.05, cy, 5.6, 1.0, m, size=18, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 11. At a glance ----------
s = new_slide("Part II · Course Design", "The course at a glance",
              notes="MGMT-205, 4 credits, prerequisite MGMT-104. Meets Monday and Thursday, 10:15 to 12:20, in AB "
                    "3336. The catalog title remains Organizational Behavior; the AI theme is the course subtitle.")
stats = [("4", "credits"), ("10", "weeks, plus finals week"), ("9", "reading reflections"), ("7", "HBS cases"),
         ("3", "levels of analysis")]
for i, (num, label) in enumerate(stats):
    cx = 0.6 + i * 2.46
    text(s, cx, 2.0, 2.3, 1.2, num, size=60, font=HEAD, bold=True, color=[TEAL, CORAL, GOLD, TEAL, CORAL][i])
    text(s, cx, 3.2, 2.2, 0.8, label, size=16, color=MUTED)
card(s, 0.6, 4.4, 12.1, 2.2)
text(s, 0.95, 4.65, 11.4, 1.8, [
    [("MGMT-205 Organizational Behavior", {"bold": True}), ("  ·  Prerequisite MGMT-104  ·  In person", {})],
    "Monday and Thursday, 10:15 A.M.–12:20 P.M.  ·  AB 3336",
    [("Textbook: ", {"bold": True}), ("Hitt, Miller, Colella, and Triana, ", {}),
     ("Organizational Behavior", {"italic": True}), (" (5th ed.), Wiley", {})],
], size=17, space_after=10)

# ---------- 12. CLOs ----------
s = new_slide("Part II · Course Design", "Course learning outcomes",
              notes="Six outcomes map to Kettering's central outcomes: Knowledge (CLO 1, 3), Reasoning (CLO 2, 3), "
                    "Ethics (CLO 4, 5), and Communication and Teamwork (CLO 6).")
clos = [("Explain core OB concepts at the individual, group, and organizational levels.", "Understand"),
        ("Apply OB theories to analyze how people and AI jointly develop knowledge and make decisions.", "Apply / Analyze"),
        ("Analyze how human–AI collaboration supports, challenges, or changes classic OB assumptions.", "Analyze"),
        ("Evaluate fairness, inclusion, accountability, and well-being implications of AI at work.", "Evaluate"),
        ("Use AI responsibly as a learning partner and reflect on how it influences your judgment.", "Apply / Evaluate"),
        ("Communicate evidence-based recommendations for human–AI work design as a team.", "Create")]
for i, (body, bloom) in enumerate(clos):
    cx, cy = 0.6 + (i % 2) * 6.15, 1.95 + (i // 2) * 1.6
    card(s, cx, cy, 5.95, 1.4)
    badge(s, cx + 0.25, cy + 0.38, 0.65, [TEAL, CORAL][i % 2], f"{i + 1}", size=18)
    text(s, cx + 1.1, cy + 0.15, 4.65, 0.85, body, size=15)
    text(s, cx + 1.1, cy + 0.98, 4.65, 0.35, bloom.upper(), size=11, bold=True, color=MUTED)

# ---------- 13. Progression ----------
s = new_slide("Part II · Course Design", "Moving up the levels of analysis",
              notes="The textbook is taught in this level-based order rather than chapter order. Each level adds "
                    "complexity to the question of how humans and AI know and decide together.")
phases = [("Week 1", "Foundations", "What does OB assume about who knows and who decides?", DARK),
          ("Weeks 2–4", "The Individual", "How does AI change trust, learning, and motivation?", TEAL),
          ("Weeks 5–7", "The Group", "How does AI change leadership, group decisions, and teams?", CORAL),
          ("Weeks 8–9", "The Organization", "How does AI change fairness, culture, and learning?", GOLD),
          ("Weeks 10–11", "Integration", "What would an OB of human–AI joint agency look like?", DARK)]
shape(s, MSO_SHAPE.RECTANGLE, 0.9, 3.02, 11.5, 0.06, LIGHT_TEAL)
for i, (wk, head, q, col) in enumerate(phases):
    cx = 0.6 + i * 2.46
    text(s, cx, 2.0, 2.3, 0.4, wk.upper(), size=12, bold=True, color=MUTED)
    badge(s, cx + 0.05, 2.7, 0.7, col, str(i + 1), size=16)
    text(s, cx, 3.7, 2.3, 0.5, head, size=19, font=HEAD, bold=True, color=col)
    text(s, cx, 4.3, 2.2, 1.8, q, size=15)

# ---------- 14. Grading ----------
s = new_slide("Part II · Course Design", "How the grade is earned",
              notes="Sixty percent of the grade now comes from work that develops over time: the reflections and "
                    "the Work Study. Exams remain, in person and without AI, to ensure everyone masters the core "
                    "concepts independently.")
cd = CategoryChartData()
cd.categories = ["Human–AI Work Study", "Examinations", "Reading reflections", "Participation"]
cd.add_series("Weight", (0.40, 0.30, 0.20, 0.10))
gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.6), Inches(1.8), Inches(6.4), Inches(5.0), cd)
ch = gf.chart
ch.has_title = False
ch.has_legend = False
plot = ch.plots[0]
plot.has_data_labels = True
dl = plot.data_labels
dl.show_value = True
dl.number_format = "0%"
dl.number_format_is_linked = False
dl.font.size = Pt(16)
dl.font.bold = True
dl.font.color.rgb = WHITE
point_colors = [CORAL, DARK, TEAL, MUTED]
for i, col in enumerate(point_colors):
    pt = plot.series[0].points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
    pt.format.line.color.rgb = WHITE
groups = [("40%", "Human–AI Work Study", "Milestones 10% · Final report 20% · Presentation 10%", CORAL),
          ("30%", "Examinations", "Midterm 15% · Final 15% · in person, no AI", DARK),
          ("20%", "Reading reflections", "Nine reflections, best 8 count", TEAL),
          ("10%", "Participation", "Graded on case discussion days", MUTED)]
for i, (pct, head, body, col) in enumerate(groups):
    cy = 1.95 + i * 1.2
    text(s, 7.5, cy, 1.5, 0.8, pct, size=34, font=HEAD, bold=True, color=col)
    text(s, 9.1, cy + 0.02, 3.7, 0.45, head, size=17, bold=True)
    text(s, 9.1, cy + 0.45, 3.7, 0.5, body, size=13, color=MUTED)

# ---------- 15. Scholarly dialogue ----------
s = new_slide("Part III · Intellectual Foundations", "Classic and contemporary, week by week",
              notes="Each pairing sets a foundational OB study against recent evidence on AI at work, mostly from the "
                    "Journal of Organizational Behavior, Academy of Management Journal, Administrative Science "
                    "Quarterly, Journal of Management Studies, and Personnel Psychology.")
table(s, 0.6, 1.85, 12.1, [0.8, 4.1, 3.6, 3.6], [
    ["Week", "Theme", "Classic OB", "Contemporary AI"],
    ["2", "Personality, attitudes, and trust in AI", "Mayer, Davis & Schoorman (1995)", "Vuori et al. (2025)"],
    ["3", "Learning and perception with algorithms", "Nonaka (1994)", "Anthony (2021)"],
    ["4", "Motivation and work design", "Hackman & Oldham (1976)", "Schulz et al. (2025)"],
    ["5", "Leadership with AI", "Graen & Uhl-Bien (1995)", "Liu et al. (2026)"],
    ["6", "Group decision making", "Stasser & Titus (1985)", "Zercher et al. (2025)"],
    ["7", "Teams, trust, and psychological safety", "Edmondson (1999)", "Erengin et al. (2025)"],
    ["8", "Diversity, inclusion, and algorithmic fairness", "Ely & Thomas (2001)", "van den Broek et al. (2025)"],
    ["9", "Culture, expertise, and organizational learning", "March (1991)", "Faulconbridge et al. (2023)"],
    ["10", "Toward an OB of human–AI joint agency", "Murray, Rhymer & Sirmon (2021)", "Stelmaszak et al. (2025)"],
], size=14, row_h=0.48)

# ---------- 16. Reading strategy ----------
s = new_slide("Part III · Intellectual Foundations", "Four questions for every pair of readings",
              notes="Undergraduates are not expected to master every methodological detail. They are expected to find "
                    "the central argument, the key evidence, and the underlying assumptions.")
qs = [("What does the classic theory explain?", "Find its central claim about behavior at work.", DARK),
      ("What does it assume about who knows and who decides?", "This is where AI puts the most pressure on theory.", TEAL),
      ("What does the AI study find?", "Which assumption does it support, challenge, or change?", CORAL),
      ("How would you revise the theory?", "For workplaces where people and AI know and decide together.", GOLD)]
for i, (q, sub, col) in enumerate(qs):
    cx, cy = 0.6 + (i % 2) * 6.15, 1.95 + (i // 2) * 2.4
    card(s, cx, cy, 5.95, 2.15)
    badge(s, cx + 0.3, cy + 0.3, 0.7, col, f"Q{i + 1}", size=16)
    text(s, cx + 1.25, cy + 0.3, 4.45, 0.95, q, size=19, bold=True)
    text(s, cx + 1.25, cy + 1.3, 4.45, 0.7, sub, size=15, color=MUTED)

# ---------- 17. JKD framework ----------
s = new_slide("Part IV · The JKD Framework", "The Joint Knowing and Deciding (JKD) Framework",
              notes="Seven questions in three clusters. Every reading, reflection, case, and project milestone maps "
                    "to one or more of them. Students should be able to ask all seven about any human–AI work setting.")
clusters = [("KNOWING", TEAL, [("1", "Who knows?", "Sources of knowledge"),
                               ("2", "How is knowledge built?", "Co-creation and validation")]),
            ("DECIDING", CORAL, [("3", "Who decides?", "Decision rights"),
                                 ("4", "Who is accountable?", "Authority and responsibility")]),
            ("CONSEQUENCES", GOLD, [("5", "How do people experience it?", "Trust, motivation, stress"),
                                    ("6", "Is it fair?", "Fairness, diversity, inclusion"),
                                    ("7", "Does the organization learn?", "Culture, expertise, adaptation")])]
for i, (name, col, items) in enumerate(clusters):
    cx = 0.6 + i * 4.1
    card(s, cx, 1.85, 3.85, 4.55)
    text(s, cx + 0.3, 2.05, 3.3, 0.4, name, size=15, bold=True, color=col)
    for j, (n, q, f) in enumerate(items):
        cy = 2.6 + j * 1.25
        badge(s, cx + 0.3, cy, 0.55, col, n, size=15)
        text(s, cx + 1.0, cy - 0.05, 2.7, 0.5, q, size=16, bold=True)
        text(s, cx + 1.0, cy + (0.72 if q.startswith("How do people") else 0.42), 2.7, 0.4, f, size=13, color=MUTED)
text(s, 0.6, 6.6, 12.1, 0.45, "Asked at three levels:  Individual  ·  Group  ·  Organization", size=17, bold=True,
     color=DARK, align=PP_ALIGN.CENTER)

# ---------- 18. Levels matrix ----------
s = new_slide("Part IV · The JKD Framework", "Seven questions, three levels",
              notes="Use this as a reference. In the Work Study, teams apply all seven questions to their chosen "
                    "organization, one level at a time.")
table(s, 0.6, 1.8, 12.1, [3.1, 3.0, 3.0, 3.0], [
    ["JKD question", "Individual", "Group", "Organization"],
    ["1. Who knows?", "What do I know vs. the AI?", "What does the team know with AI?", "Where does knowledge reside?"],
    ["2. How is knowledge built?", "Am I learning or outsourcing?", "Does AI help pool knowledge?", "Is AI knowledge captured and checked?"],
    ["3. Who decides?", "When do I follow AI advice?", "How do teams decide with AI?", "What decision rights go to AI?"],
    ["4. Who is accountable?", "Am I responsible for AI-shaped choices?", "Who owns team decisions?", "What governance assigns responsibility?"],
    ["5. How do people experience it?", "Trust, motivation, stress", "Team trust and cohesion", "Climate and well-being"],
    ["6. Is it fair?", "Am I treated fairly by AI?", "Are members treated equally?", "Are systems fair across groups?"],
    ["7. Does the organization learn?", "Do I keep building expertise?", "Can the team question AI?", "Does the culture support adaptation?"],
], size=13, row_h=0.6, first_col_bold=True)

# ---------- 19. JKD across the term ----------
s = new_slide("Part IV · The JKD Framework", "Where each question takes center stage",
              notes="Weeks 1 and 10 address all seven questions: Week 1 as an introduction, Week 10 as the synthesis. "
                    "In between, each week foregrounds one to three questions.")
week_q = {1: range(1, 8), 2: [5], 3: [1, 2], 4: [5], 5: [3, 4], 6: [2, 3], 7: [4, 5, 7], 8: [6], 9: [1, 7],
          10: range(1, 8)}
qlabels = ["1  Who knows?", "2  How is knowledge built?", "3  Who decides?", "4  Who is accountable?",
           "5  How do people experience it?", "6  Is it fair?", "7  Does the organization learn?"]
qcol = [TEAL, TEAL, CORAL, CORAL, GOLD, GOLD, GOLD]
gx, gy, cw, rh = 4.55, 2.2, 0.8, 0.6
for w in range(1, 11):
    text(s, gx + (w - 1) * cw, 1.75, cw, 0.35, f"Wk {w}", size=12, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
for q in range(7):
    text(s, 0.6, gy + q * rh, 3.8, rh, qlabels[q], size=14, bold=True, color=qcol[q], anchor=MSO_ANCHOR.MIDDLE)
    for w in range(1, 11):
        on = (q + 1) in week_q[w]
        d = 0.36 if on else 0.14
        cx = gx + (w - 1) * cw + (cw - d) / 2
        cy = gy + q * rh + (rh - d) / 2
        shape(s, MSO_SHAPE.OVAL, cx, cy, d, d, qcol[q] if on else TINT)
text(s, 0.6, 6.6, 12.1, 0.45, "Large dots mark each week's primary JKD questions.", size=13, italic=True,
     color=MUTED)

# ---------- 20. OB Theory Update ----------
s = new_slide("Part IV · The JKD Framework", "The OB Theory Update",
              notes="Each weekly reflection closes with an OB Theory Update. The strong example names the original "
                    "assumption, cites evidence, proposes a specific change, and explains why it fits the evidence.")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.85, 12.1, 1.35, DARK, radius=0.08)
text(s, 0.95, 1.95, 11.4, 1.15, "“If [this week's classic theory] were rewritten for workplaces where people and AI "
     "build knowledge and make decisions together, I would change… because [evidence].”", size=19, italic=True,
     color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
card(s, 0.6, 3.5, 4.3, 3.15)
text(s, 0.95, 3.7, 3.6, 0.4, "WEAK", size=14, bold=True, color=CORAL)
text(s, 0.95, 4.2, 3.6, 2.3, "“Trust theory needs to change because people trust AI differently.”", size=17,
     italic=True)
card(s, 5.2, 3.5, 7.5, 3.15)
text(s, 5.55, 3.7, 6.8, 0.4, "STRONG", size=14, bold=True, color=TEAL)
text(s, 5.55, 4.15, 6.8, 2.4, "“Mayer, Davis, and Schoorman define trust through ability, benevolence, and integrity, "
     "assuming the trustee has intentions. Vuori et al. (2025) show emotional trust in AI still shapes use. I would "
     "replace benevolence with ‘perceived alignment’ — the belief the system was designed in the user's interest.”",
     size=15, italic=True)

# ---------- 21. Closing ----------
s = new_slide(dark=True, notes="Close by reading the seven questions aloud. These are the practical core of "
              "organizational behavior in the age of AI, and students will use them in every week of the course.")
motif(s, 0.8, 0.9, 0.8)
text(s, 0.8, 2.0, 6.0, 2.6, "Seven questions to ask of any workplace where people and AI work together",
     size=32, font=HEAD, bold=True, color=WHITE)
text(s, 0.8, 5.2, 5.8, 1.2, "Humans and AI — knowing and deciding together.", size=20, italic=True,
     color=LIGHT_TEAL)
closing_q = ["Who knows?", "How is knowledge built?", "Who decides?", "Who is accountable?",
             "How do people experience it?", "Is it fair?", "Does the organization learn?"]
for i, q in enumerate(closing_q):
    cy = 0.9 + i * 0.85
    badge(s, 7.4, cy, 0.55, qcol[i], str(i + 1), size=15)
    text(s, 8.2, cy, 4.6, 0.55, q, size=20, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
