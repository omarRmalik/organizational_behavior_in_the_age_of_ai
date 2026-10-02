"""Build the MGMT-205 Organizational Behavior Fall 2026 syllabus as a Word document."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import sys
OUT = sys.argv[1] if len(sys.argv) > 1 else "mgt_205_fall_2026_malik.docx"
TERM = "Fall 2026"
FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "D9D9D9"
BAND_FILL = "F2F2F2"
REPO = "https://github.com/omarRmalik/organizational_behavior_in_the_age_of_ai"

doc = Document()

# ---------- page setup & base styles ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.orientation = WD_ORIENT.PORTRAIT
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.color.rgb = BLACK
normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, before in (("Title", 20, 0), ("Heading 1", 15, 16), ("Heading 2", 12.5, 10)):
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = BLACK
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
    rpr = st.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), FONT)
    fonts.attrib.pop(qn("w:asciiTheme"), None)
    fonts.attrib.pop(qn("w:hAnsiTheme"), None)

# The default template's Title style has a colored bottom border; drop it.
for bdr in doc.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
    bdr.getparent().remove(bdr)


def add_field(run, instr):
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


# header / footer
hp = sec.header.paragraphs[0]
hp.text = f"MGMT-205 — Organizational Behavior in the Age of AI — {TERM}"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size = Pt(9)
    r.font.name = FONT
    r.font.color.rgb = BLACK

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for piece in ("Page ", "PAGE", " of ", "NUMPAGES"):
    r = fp.add_run()
    r.font.size = Pt(9)
    r.font.name = FONT
    r.font.color.rgb = BLACK
    if piece.isupper():
        add_field(r, piece)
    else:
        r.text = piece


# ---------- helpers ----------
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def table(headers, rows, widths, center_cols=(), first_col_bold=False, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        cell_text(c, h, bold=True, size=size,
                  align=WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else None)
        shade(c, HEADER_FILL)
    for ri, row in enumerate(rows):
        r = t.add_row()
        no_split(r)
        for i, val in enumerate(row):
            c = r.cells[i]
            cell_text(c, val, bold=(first_col_bold and i == 0), size=size,
                      align=WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else None)
            if ri % 2 == 1:
                shade(c, BAND_FILL)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def kv_table(pairs, key_w=2.1, val_w=4.4):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.autofit = False
    for k, v in pairs:
        r = t.add_row()
        no_split(r)
        cell_text(r.cells[0], k, bold=True, size=10.5)
        shade(r.cells[0], BAND_FILL)
        cell_text(r.cells[1], v, size=10.5)
        r.cells[0].width, r.cells[1].width = Inches(key_w), Inches(val_w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def para(text, italic=False):
    p = doc.add_paragraph()
    p.add_run(text).italic = italic
    return p


def rich(parts):
    """parts: list of (text, style) where style in {'', 'b', 'i'}."""
    p = doc.add_paragraph()
    for text, style in parts:
        r = p.add_run(text)
        r.bold = "b" in style
        r.italic = "i" in style
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def ref(parts):
    """Hanging-indent APA reference. parts as in rich()."""
    p = rich(parts)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    return p


def add_hyperlink(paragraph, text, url, size=None):
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                   is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), FONT)
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)




from urllib.parse import quote
import re


def doi(d):
    return "https://doi.org/" + quote(d, safe="/:()")


# Every reading links to its DOI; open-access ones are free to all, the rest open through library access.
READING_LINKS = {
    "Mayer, Davis & Schoorman (1995)": doi("10.5465/amr.1995.9508080335"),
    "Vuori et al. (2025)": doi("10.1111/joms.13177"),
    "Nonaka (1994)": doi("10.1287/orsc.5.1.14"),
    "Anthony (2021)": doi("10.1177/00018392211016755"),
    "Hackman & Oldham (1976)": doi("10.1016/0030-5073(76)90016-7"),
    "Schulz et al. (2025)": doi("10.1111/joms.70004"),
    "Graen & Uhl-Bien (1995)": doi("10.1016/1048-9843(95)90036-5"),
    "Liu et al. (2026)": doi("10.1002/job.70070"),
    "Stasser & Titus (1985)": doi("10.1037/0022-3514.48.6.1467"),
    "Zercher et al. (2025)": doi("10.1002/job.2898"),
    "Edmondson (1999)": doi("10.2307/2666999"),
    "Erengin et al. (2025)": doi("10.1002/job.2857"),
    "Ely & Thomas (2001)": doi("10.2307/2667087"),
    "van den Broek et al. (2025)": doi("10.1111/joms.13276"),
    "March (1991)": doi("10.1287/orsc.2.1.71"),
    "Faulconbridge et al. (2023)": doi("10.1111/joms.12936"),
    "Murray, Rhymer & Sirmon (2021)": doi("10.5465/amr.2019.0186"),
    "Stelmaszak et al. (2025)": doi("10.1111/joms.70003"),
}
_READING_RE = re.compile("|".join(re.escape(k) for k in sorted(READING_LINKS, key=len, reverse=True)))


def linked_text(paragraph, text, size=None):
    """Write text into paragraph, turning known reading names into hyperlinks."""
    pos = 0
    for m in _READING_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            if size:
                r.font.size = Pt(size)
        add_hyperlink(paragraph, m.group(), READING_LINKS[m.group()], size=size)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        if size:
            r.font.size = Pt(size)


def relink_column(tbl, col, size):
    for row in tbl.rows[1:]:
        cell_p = row.cells[col].paragraphs[0]
        text = cell_p.text
        for r in list(cell_p.runs):
            r._r.getparent().remove(r._r)
        linked_text(cell_p, text, size=size)


# ---------- title ----------
doc.add_paragraph("Organizational Behavior in the Age of Artificial Intelligence", style="Title")
sub = doc.add_paragraph()
r = sub.add_run(f"MGMT-205-01  •  {TERM}  •  School of Management, Kettering University")
r.font.size = Pt(12)
r.font.color.rgb = BLACK
r.bold = True

# ---------- course & instructor info ----------
h1("Course Information")
kv_table([
    ("Course Number", "MGMT-205-01"),
    ("Course Title", "Organizational Behavior"),
    ("Course Theme", "Organizational Behavior in the Age of Artificial Intelligence"),
    ("Credit Hours", "4"),
    ("Term", TERM),
    ("Meeting Days/Times", "Monday and Thursday, 10:15 A.M.–12:20 P.M."),
    ("Location", "AB 3336"),
    ("Delivery Format", "In-Person"),
    ("Prerequisite", "MGMT-104"),
])

h1("Instructor Information")
kv_table([
    ("Instructor", "Omar R. Malik, Ph.D."),
    ("Office", "AB 4103"),
    ("Email", "omalik@kettering.edu"),
    ("Phone", "(616) 808-9866"),
    ("Office Hours", "Monday and Thursday, 12:20–1:20 P.M."),
    ("Preferred Communication Method", "Email"),
])

# ---------- description ----------
h1("Course Description")
para("The central task in organizations is working with people. Increasingly, people work alongside artificial "
     "intelligence that searches, summarizes, predicts, recommends, evaluates, and decides. The organizational "
     "behavior discipline examines the behavior of individuals, groups, and organizations as it relates to desirable "
     "organizational outcomes such as motivation, productivity, ethical behavior, and social responsibility.")
para("In this course we survey research-based findings in organizational behavior—personality and attitudes, "
     "learning and perception, motivation, leadership, group decision making, teams, diversity, and culture—and "
     "examine how each changes when humans and AI jointly develop knowledge and make decisions. Students pair the "
     "textbook with classic and contemporary research, analyze Harvard Business School cases, "
     " and study a real organization's human–AI work practices in a team project.")

h2("The Central Question")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.4)
p.add_run("How does organizational behavior change when humans and AI jointly develop knowledge and make "
          "decisions?").bold = True

h2("The Joint Knowing and Deciding (JKD) Framework")
para("Every topic in the course is examined through seven questions, asked at the individual, group, and "
     "organizational levels:")
table(["Cluster", "Question", "Focus"], [
    ["Knowing", "1. Who knows?", "Sources of knowledge"],
    ["Knowing", "2. How is knowledge built?", "Co-creation and validation"],
    ["Deciding", "3. Who decides?", "Decision rights"],
    ["Deciding", "4. Who is accountable?", "Authority and responsibility"],
    ["Consequences", "5. How do people experience it?", "Trust, motivation, stress, identity"],
    ["Consequences", "6. Is it fair?", "Fairness, diversity, inclusion"],
    ["Consequences", "7. Does the organization learn?", "Culture, expertise, adaptation"],
], widths=[1.5, 2.6, 2.4])

# ---------- outcomes ----------
h1("Course Learning Outcomes (CLOs)")
para("The learning outcomes of this course relate to Kettering University's central learning outcomes of "
     "Knowledge, Reasoning, and Ethics. We also develop communication and teamwork skills throughout the course. "
     "Upon successful completion of this course, students will be able to:")
table(["CLO", "Learning Outcome", "Bloom's Level"], [
    ["CLO 1", "Explain core organizational behavior concepts at the individual, group, and organizational levels.",
     "Understand"],
    ["CLO 2", "Apply organizational behavior theories to analyze how people and AI jointly develop knowledge and make "
              "decisions in real organizations.", "Apply / Analyze"],
    ["CLO 3", "Analyze how human–AI collaboration supports, challenges, or changes the assumptions of classic "
              "organizational behavior theories.", "Analyze"],
    ["CLO 4", "Evaluate the fairness, inclusion, accountability, and well-being implications of AI in workplace "
              "decisions.", "Evaluate"],
    ["CLO 5", "Use AI tools responsibly as a partner in learning and decision making, and reflect critically on how "
              "AI influences one's own judgment.", "Apply / Evaluate"],
    ["CLO 6", "Communicate evidence-based recommendations for human–AI work design, in writing and orally, as part "
              "of a team.", "Create"],
], widths=[0.8, 4.4, 1.3], center_cols=(0, 2), first_col_bold=True)

h2("Alignment with University Learning Outcomes")
table(["University Outcome", "Course Learning Outcome(s)"], [
    ["Knowledge", "CLO 1, CLO 3"],
    ["Reasoning", "CLO 2, CLO 3"],
    ["Ethics", "CLO 4, CLO 5"],
    ["Communication and Teamwork", "CLO 6"],
], widths=[3.0, 3.5])

# ---------- materials ----------
h1("Required Course Materials")
h2("Textbook")
para("Hitt, Miller, Colella, and Triana. Organizational Behavior (5th ed.). Wiley. ISBN 978-1-119-39173-9.")
h2("Course Packet")
para("Harvard Business School Publishing course packet. The Fall 2026 packet link is posted on the LMS.")
h2("Research Articles")
para("Beginning in Week 2, each week pairs a classic organizational behavior article with a contemporary article on AI at work, mainly "
     "from the Journal of Organizational Behavior, Academy of Management Journal, Administrative Science Quarterly, "
     "Journal of Management Studies, and Personnel Psychology. Reading names in the course schedule are links. "
     "Many of the AI articles are open access; the others open through Kettering library access.")
h2("AI Tools")
para("You may use a university-approved AI assistant (e.g., Claude, ChatGPT, Copilot, or Gemini) as a learning "
     "partner, following the AI policy below. No programming experience is required.")

# ---------- schedule ----------
h1("Weekly Class Schedule")
rich([("For complete weekly module details, see ", ""),
      ("Part V – Course Roadmap and Weekly Learning Modules", "i"), (".", "")])
schedule = table(["Week", "Topic", "Textbook", "Readings", "HBS Case", "Due"], [
    ["1", "OB when humans and AI think together", "Ch. 1: A strategic approach to OB",
     "—", "—", "Team proposal"],
    ["2", "Personality, attitudes, and trust in AI", "Ch. 5: Personality, intelligence, attitudes, and emotions",
     "Mayer, Davis & Schoorman (1995); Vuori et al. (2025)",
     "Microsoft Customer and Partner Solutions: The Deployment of Copilot (A)", "Organization approved"],
    ["3", "Learning and perception with algorithms", "Ch. 4: Learning and perception",
     "Nonaka (1994); Anthony (2021)", "—", "Milestone 1"],
    ["4", "Motivation and work design", "Ch. 6: Work motivation",
     "Hackman & Oldham (1976); Schulz et al. (2025)", "Trouble at Tessei", "—"],
    ["5", "Leadership with AI", "Leadership chapter",
     "Graen & Uhl-Bien (1995); Liu et al. (2026)", "Google's Project Oxygen", "Midterm exam"],
    ["6", "Group decision making", "Ch. 10: Decision making by individuals and groups",
     "Stasser & Titus (1985); Zercher et al. (2025)", "JPMorganChase: Leadership in the Age of GenAI",
     "Milestone 2"],
    ["7", "Teams, trust, and psychological safety", "Groups and teams chapter",
     "Edmondson (1999); Erengin et al. (2025)", "Governing OpenAI", "—"],
    ["8", "Diversity, inclusion, and algorithmic fairness", "Ch. 2: Organizational diversity",
     "Ely & Thomas (2001); van den Broek et al. (2025)", "Dessa: Growing a Diverse and Inclusive AI Company",
     "Milestone 3"],
    ["9", "Culture, expertise, and organizational learning", "Organizational culture chapter",
     "March (1991); Faulconbridge et al. (2023)",
     "Culture Transformation at Microsoft: From \"Know It All\" to \"Learn It All\"", "Draft report"],
    ["10", "Toward an OB of human–AI joint agency", "Review",
     "Murray, Rhymer & Sirmon (2021); Stelmaszak et al. (2025)", "—", "Team presentations"],
    ["11", "Finals week", "—", "—", "—", "Final report; final exam"],
], widths=[0.45, 1.2, 1.2, 1.45, 1.25, 0.95], center_cols=(0,), size=9)
relink_column(schedule, 3, 9)
para("Weekly reading reflections are due in Weeks 2 through 10 (see Course Components).")

# ---------- components ----------
h1("Course Components")
h2("Attendance")
para("I expect you to attend every class. If an exceptional circumstance prevents you from attending, please let me "
     "know before class, and I will work with you to help cover the content you missed. There is no grade for "
     "attendance, but case classes require your participation, so attendance is critical.")

h2("Class Participation")
para("Participation is graded on case discussion days on a 0–3 scale per class. I will share scores with you after "
     "each case class.")
table(["Score", "Level", "Description"], [
    ["0", "No substantive comment", "Does not add to the discussion or repeats others; superficial, irrelevant, or "
                                    "random; not listening, interrupting, or arguing in a hostile manner."],
    ["1–2", "Straightforward comment", "Adds to understanding; responds adequately to questions; uses concepts and "
                                       "frameworks to link the discussion; engages constructively with classmates."],
    ["3", "Insightful comment", "Significantly improves understanding; shows deep and complex thinking; connects the "
                                "case to OB theory and the course's central question."],
], widths=[0.7, 1.7, 4.1], center_cols=(0,))
para("If you have problems speaking up in group discussions, please speak to me early in the term. Do not wait.")

h2("Weekly Reading Reflections")
para("Beginning in Week 2, each week pairs a classic organizational behavior article with a contemporary article on "
     "AI at work. After each week's readings you write a 400–600 word reflection with three parts: (1) a reading "
     "connection—what the classic theory explains, what it assumes about who knows and who decides, and what the AI "
     "study shows; (2) an OB Theory Update; and (3) an AI disclosure. The Theory Update responds to the prompt:")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.4)
p.add_run("\"If this OB theory were rewritten for workplaces where people and AI build knowledge and make decisions "
          "together, I would change…\"").italic = True
para("There are nine reflections (Weeks 2–10). Each is scored out of 10; your best 8 of 9 count.")

h2("Human–AI Work Study (Team Project)")
para("Teams of three study how a real organization, or a department within it, uses AI in knowledge work and "
     "decisions, and how that affects its people. Teams apply the JKD Framework one level at a time and recommend "
     "changes to job design, training, decision rights, trust calibration, and fairness safeguards.")
bullets([
    ("Milestones (10%): ", "three 2-page memos—individual level (Week 3), group level (Week 6), and organization "
                          "level (Week 8)."),
    ("Final report (20%): ", "no more than 8 pages, due in finals week, with a draft due in Week 9."),
    ("Presentation (10%): ", "12 minutes plus questions, in Week 10."),
    ("Peer evaluation: ", "when peer evaluations show a clear pattern of unequal contribution, an individual's project "
                          "grade may be adjusted by up to one letter grade."),
])
rich([("See ", ""), ("Part VII – The Human–AI Work Study", "i"), (" for full details.", "")])

h2("Examinations")
para("The midterm (Week 5) and final (finals week) exams use multiple-choice and short-answer questions on OB "
     "concepts and their application to workplace situations, including situations involving AI. Exams are taken in "
     "person without AI tools. The final is cumulative with emphasis on Weeks 5–10.")

# ---------- grading ----------
h1("Grading Components")
t = table(["Assessment", "Weight"], [
    ["Class Participation (case discussions)", "10%"],
    ["Weekly Reading Reflections", "20%"],
    ["Midterm Examination", "15%"],
    ["Human–AI Work Study — Milestones", "10%"],
    ["Human–AI Work Study — Final Report", "20%"],
    ["Human–AI Work Study — Presentation", "10%"],
    ["Final Examination", "15%"],
    ["Total", "100%"],
], widths=[4.5, 1.2], center_cols=(1,))
for c in t.rows[-1].cells:
    for r in c.paragraphs[0].runs:
        r.bold = True

h2("Grading Scale")
table(["Grade", "Range", "Grade", "Range", "Grade", "Range", "Grade", "Range"], [
    ["A", "93–100", "B+", "87–89", "C+", "77–79", "D+", "67–69"],
    ["A−", "90–92", "B", "83–86", "C", "73–76", "D", "63–66"],
    ["", "", "B−", "80–82", "C−", "70–72", "F", "Below 63"],
], widths=[0.8] * 8, center_cols=tuple(range(8)))

h2("Course Learning Outcome Assessment Matrix")
matrix = {
    "Class Participation": {1, 2, 4},
    "Weekly Reading Reflections": {3, 5},
    "Midterm Examination": {1, 2},
    "Work Study Milestones": {2, 3},
    "Work Study Final Report": {2, 3, 4, 6},
    "Work Study Presentation": {4, 6},
    "Final Examination": {1, 2, 3},
}
table(["Assessment"] + [f"CLO {i}" for i in range(1, 7)],
      [[name] + ["X" if i in clos else "" for i in range(1, 7)] for name, clos in matrix.items()],
      widths=[2.3] + [0.7] * 6, center_cols=tuple(range(1, 7)), size=9)

# ---------- dates ----------
h1("Important Dates")
table(["Week", "Activity"], [
    ["Week 1", "Teams formed; preliminary organization proposal"],
    ["Week 2", "Project organization approved"],
    ["Week 3", "Milestone 1: individual-level memo"],
    ["Week 5", "Midterm examination"],
    ["Week 6", "Milestone 2: group-level memo"],
    ["Week 8", "Milestone 3: organization-level memo"],
    ["Week 9", "Draft report"],
    ["Week 10", "Team presentations"],
    ["Finals week", "Final report and peer evaluation; final examination"],
], widths=[1.3, 5.2], first_col_bold=True)
para("Weekly reading reflections (Weeks 2–10) are due before the start of the following week's first class. Exact dates and submission "
     "times are posted on the LMS.")

# ---------- policies ----------
h1("Course Policies")
h2("Artificial Intelligence (AI) Policy")
para("In this course AI is both a subject of study and a learning partner. You are welcome to use AI to explain "
     "concepts, generate practice questions, brainstorm, and improve your "
     "writing. You remain responsible for everything you submit: verify AI output against the textbook and readings, "
     "cite only sources you have read, and include an AI disclosure statement in every reflection, memo, report, and "
     "presentation. Honest disclosure is never penalized.")
para("Never enter personal information about identifiable people (including classmates, co-op colleagues, or "
     "interviewees) or confidential organizational information into AI tools. AI tools may not be used during exams.")
rich([("For the complete policy, see ", ""),
      ("Part IX – Responsible AI, Academic Integrity, and Professional Conduct", "i"), (".", "")])

h2("Academic Integrity")
para("All work submitted in this course must be your own. Fabricating citations, presenting AI-generated or others' "
     "work as your own, misrepresenting sources, or collaborating on individual work without permission constitutes "
     "academic dishonesty and will be addressed according to university policy.")

h2("Late Work")
para("Reflections and milestones submitted late lose 10% per day, up to three days, after which they receive no "
     "credit unless an extension was arranged in advance. The final report cannot be accepted after the end of "
     "finals week without an approved incomplete.")

h2("Communication")
para("Please use email to contact me. I typically respond within 48 hours on business days. Course announcements "
     "are posted on the LMS.")

h2("Accessibility and Accommodations")
para("Students who need accommodations because of a disability should contact the university's disability services "
     "office and provide documentation as early as possible. I am committed to equitable access to all course "
     "materials and activities.")

h2("University Syllabus")
para("The university syllabus, including policies on academic integrity, disability services, Title IX, student "
     "conduct, and grade appeals, is posted on Blackboard and is incorporated by reference.")

# ---------- documentation ----------
h1("Full Course Documentation")
p = doc.add_paragraph("This syllabus is supported by the following course documents, available at ")
add_hyperlink(p, REPO, REPO)
p.add_run(". Click a document name to open it.")
DOC_FILES = ["part_0_course_philosophy", "part_1_course_foundations", "part_2_course_design",
             "part_3_intellectual_foundations", "part_4_the_jkd_framework", "part_5_weekly_learning_modules",
             "part_6_reading_reflections", "part_7_human_ai_work_study", "part_8_assessment_standards",
             "part_9_policies", "part_10_reading_list"]
docs_table = table(["Document", "Description"], [
    ["Part 0 – Course Philosophy", "Why this course exists and the question it asks"],
    ["Part I – Course Foundations", "Pedagogical model and expected student mindset"],
    ["Part II – Course Design", "Learning outcomes, assurance of learning, progression"],
    ["Part III – Intellectual Foundations", "The classic OB canon and the AI-at-work canon"],
    ["Part IV – The JKD Framework", "Joint Knowing and Deciding: the organizing framework"],
    ["Part V – Weekly Learning Modules", "Week-by-week course roadmap"],
    ["Part VI – Weekly Reading Reflections", "Reflection design and the OB Theory Update"],
    ["Part VII – The Human–AI Work Study", "Team project, milestones, and deliverables"],
    ["Part VIII – Assessment Standards", "Grading, rubrics, and assurance of learning"],
    ["Part IX – Policies", "Responsible AI, academic integrity, professional conduct"],
    ["Part X – Reading List", "Complete reading list with DOIs"],
], widths=[3.0, 3.5])
for row, fname in zip(docs_table.rows[1:], DOC_FILES):
    cell_p = row.cells[0].paragraphs[0]
    label = cell_p.text
    for r in list(cell_p.runs):
        r._r.getparent().remove(r._r)
    add_hyperlink(cell_p, label, f"{REPO}/blob/main/md%20files/{fname}.md", size=10)

doc.core_properties.title = f"MGMT-205 Organizational Behavior Syllabus — {TERM}"
doc.core_properties.author = "Omar R. Malik, Ph.D."
doc.save(OUT)
print("saved", OUT)
